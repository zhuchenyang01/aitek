"""PDF / Word 文档读取：路径输入，分离文字与图片。"""
from __future__ import annotations

from pathlib import Path
from typing import Any

from rag.utils.timing import timed


def _ext(path: Path) -> str:
    return path.suffix.lower()


@timed("document_loader.load")
def load_document(file_path: str | Path, image_out_dir: str | Path) -> list[dict[str, Any]]:
    """
    读取 PDF/Word，返回有序块列表：
    {"type": "text"|"image", "content": str|None, "path": str|None, "order": int, "page": int|None}
    """
    path = Path(file_path).expanduser().resolve()
    out_dir = Path(image_out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    if not path.exists():
        raise FileNotFoundError(f"文档不存在: {path}")

    print(f"[StepLoader] 读取文档: {path}")
    if _ext(path) == ".pdf":
        blocks = _load_pdf(path, out_dir)
    elif _ext(path) in {".docx", ".doc"}:
        if _ext(path) == ".doc":
            raise ValueError("暂仅支持 .docx，请先转换为 docx")
        blocks = _load_docx(path, out_dir)
    else:
        raise ValueError(f"不支持的文件类型: {_ext(path)}，仅支持 pdf/docx")

    text_n = sum(1 for b in blocks if b["type"] == "text")
    img_n = sum(1 for b in blocks if b["type"] == "image")
    print(f"[StepLoader] 完成: 文本块={text_n}, 图片={img_n}, 总块={len(blocks)}")
    return blocks


def _load_pdf(path: Path, out_dir: Path) -> list[dict[str, Any]]:
    import fitz  # pymupdf
    from pypdf import PdfReader

    blocks: list[dict[str, Any]] = []
    order = 0

    reader = PdfReader(str(path))
    page_texts: list[str] = []
    for page in reader.pages:
        text = (page.extract_text() or "").strip()
        page_texts.append(text)

    doc = fitz.open(str(path))
    for page_index in range(len(doc)):
        page = doc[page_index]
        text = page_texts[page_index] if page_index < len(page_texts) else ""
        if text:
            blocks.append(
                {
                    "type": "text",
                    "content": text,
                    "path": None,
                    "order": order,
                    "page": page_index + 1,
                }
            )
            order += 1

        for img_i, img in enumerate(page.get_images(full=True)):
            xref = img[0]
            try:
                pix = fitz.Pixmap(doc, xref)
                if pix.n >= 5:
                    pix = fitz.Pixmap(fitz.csRGB, pix)
                img_path = out_dir / f"pdf_p{page_index + 1}_{img_i + 1}.png"
                pix.save(str(img_path))
                blocks.append(
                    {
                        "type": "image",
                        "content": None,
                        "path": str(img_path),
                        "order": order,
                        "page": page_index + 1,
                    }
                )
                order += 1
            except Exception as exc:  # noqa: BLE001
                print(f"[StepLoader] 跳过损坏图片 page={page_index + 1} img={img_i}: {exc}")

    doc.close()
    return blocks


def _load_docx(path: Path, out_dir: Path) -> list[dict[str, Any]]:
    from docx import Document

    doc = Document(str(path))
    blocks: list[dict[str, Any]] = []
    order = 0
    img_idx = 0

    for para in doc.paragraphs:
        text = (para.text or "").strip()
        if text:
            blocks.append(
                {
                    "type": "text",
                    "content": text,
                    "path": None,
                    "order": order,
                    "page": None,
                }
            )
            order += 1

    for table in doc.tables:
        rows = []
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            rows.append(" | ".join(cells))
        table_text = "\n".join(rows).strip()
        if table_text:
            blocks.append(
                {
                    "type": "text",
                    "content": table_text,
                    "path": None,
                    "order": order,
                    "page": None,
                }
            )
            order += 1

    for rel in doc.part.rels.values():
        if "image" not in rel.reltype:
            continue
        img_idx += 1
        blob = rel.target_part.blob
        ctype = getattr(rel.target_part, "content_type", "") or ""
        ext = ".png"
        if "jpeg" in ctype or "jpg" in ctype:
            ext = ".jpg"
        elif "gif" in ctype:
            ext = ".gif"
        elif "bmp" in ctype:
            ext = ".bmp"
        img_path = out_dir / f"docx_img_{img_idx}{ext}"
        img_path.write_bytes(blob)
        blocks.append(
            {
                "type": "image",
                "content": None,
                "path": str(img_path),
                "order": order,
                "page": None,
            }
        )
        order += 1

    blocks.sort(key=lambda b: b["order"])
    return blocks
