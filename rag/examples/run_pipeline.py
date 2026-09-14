#!/usr/bin/env python3
"""端到端演示：Step1 -> Step9，每步清晰打印。"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

# 保证可从任意目录运行：把项目根加入 PYTHONPATH
ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from rag import config  # noqa: E402
from rag.pipeline import (  # noqa: E402
    step1_parse_document,
    step2_vision_to_text,
    step3_merge_pure_text,
    step4_chunk_and_embed,
    step5_extract_features,
    step6_feature_embed_store,
    step7_associate_retrieve,
    step8_rerank,
    step9_llm_enhance,
)


def _ensure_sample_docx(path: Path) -> Path:
    """若未提供文档，生成一份示例 docx 便于演示。"""
    from docx import Document

    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading("AItek 需求示例：用户登录与项目配置", level=1)
    doc.add_paragraph(
        "模块一：用户系统。支持用户注册、登录、退出。"
        "注册需校验用户名唯一、密码不少于 6 位；登录成功后签发 Token。"
    )
    doc.add_paragraph(
        "模块二：项目配置。支持配置项的新增、编辑、删除与列表查询。"
        "字段包括 KEY、VALUE、备注、创建时间。KEY 必须唯一。"
    )
    doc.add_paragraph(
        "模块三：鉴权。除登录注册外，业务接口需携带 Bearer Token；"
        "未登录返回 401。登录注册接口可对请求体做 AES 加解密。"
    )
    doc.add_paragraph(
        "非功能：接口需可观测、关键操作有提示；前端侧边栏可导航到各测试模块。"
    )
    doc.save(str(path))
    print(f"[Demo] 已生成示例文档: {path}")
    return path


def main() -> None:
    parser = argparse.ArgumentParser(description="AItek 自研 RAG 流水线（无 LangChain）")
    parser.add_argument(
        "--doc",
        type=str,
        default="",
        help="需求文档路径（pdf/docx）。不传则自动生成示例 docx",
    )
    args = parser.parse_args()

    config.ensure_dirs()
    if args.doc:
        doc_path = Path(args.doc).expanduser().resolve()
    else:
        doc_path = _ensure_sample_docx(
            Path(__file__).resolve().parent / "sample_docs" / "sample_requirements.docx"
        )

    print("\n########## AItek RAG Pipeline 开始 ##########\n")
    blocks = step1_parse_document.run(doc_path)
    blocks2 = step2_vision_to_text.run(blocks)
    pure = step3_merge_pure_text.run(blocks2)
    step4_chunk_and_embed.run(pure)
    features = step5_extract_features.run(pure)
    rows, store = step6_feature_embed_store.run(features)
    associations = step7_associate_retrieve.run(rows, store)
    reranked = step8_rerank.run(associations)
    step9_llm_enhance.run(reranked)
    print("\n########## 全部步骤完成 ##########")
    print(f"产物目录: {config.DATA_DIR}")
    print(f"  - 纯文本: {config.PURE_TEXT_PATH}")
    print(f"  - 功能点 jsonl: {config.FEATURE_JSONL}")
    print(f"  - FAISS: {config.FAISS_INDEX_PATH}")
    print(f"  - 关联: {config.ASSOC_JSON}")
    print(f"  - 重排: {config.RERANK_JSON}")
    print(f"  - LLM增强: {config.ENHANCE_MD}")


if __name__ == "__main__":
    main()
