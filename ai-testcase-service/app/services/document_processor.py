from pathlib import Path
from typing import Any


def _ext(path: Path) -> str:
    return path.suffix.lower()


def load_document(file_path: str | Path, image_out_dir: str | Path) -> list[dict[str, Any]]:
    path = Path(file_path).expanduser().resolve()
    out_dir = Path(image_out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    if not path.exists():
        raise FileNotFoundError(f'文档不存在: {path}')

    if _ext(path) == '.pdf':
        return _load_pdf(path, out_dir)
    if _ext(path) in {'.docx', '.doc'}:
        if _ext(path) == '.doc':
            raise ValueError('暂仅支持 .docx，请先转换为 docx')
        return _load_docx(path, out_dir)
    raise ValueError(f'不支持的文件类型: {_ext(path)}，仅支持 pdf/docx')


def blocks_to_plain_text(blocks):
    parts = []
    for block in blocks or []:
        if block.get('type') == 'text' and block.get('content'):
            parts.append(block['content'])
    return '\n\n'.join(parts)


def _load_pdf(path: Path, out_dir: Path) -> list[dict[str, Any]]:
    import fitz
    from pypdf import PdfReader

    blocks: list[dict[str, Any]] = []
    order = 0
    reader = PdfReader(str(path))
    page_texts = [(page.extract_text() or '').strip() for page in reader.pages]

    doc = fitz.open(str(path))
    for page_index in range(len(doc)):
        page = doc[page_index]
        text = page_texts[page_index] if page_index < len(page_texts) else ''
        if text:
            blocks.append({'type': 'text', 'content': text, 'path': None, 'order': order, 'page': page_index + 1})
            order += 1

        for img_i, img in enumerate(page.get_images(full=True)):
            xref = img[0]
            try:
                pix = fitz.Pixmap(doc, xref)
                if pix.n >= 5:
                    pix = fitz.Pixmap(fitz.csRGB, pix)
                img_path = out_dir / f'pdf_p{page_index + 1}_{img_i + 1}.png'
                pix.save(str(img_path))
                blocks.append({'type': 'image', 'content': None, 'path': str(img_path), 'order': order, 'page': page_index + 1})
                order += 1
            except Exception:
                continue
    doc.close()
    return blocks


def _load_docx(path: Path, out_dir: Path) -> list[dict[str, Any]]:
    from docx import Document

    doc = Document(str(path))
    blocks: list[dict[str, Any]] = []
    order = 0
    img_idx = 0

    for para in doc.paragraphs:
        text = (para.text or '').strip()
        if text:
            blocks.append({'type': 'text', 'content': text, 'path': None, 'order': order, 'page': None})
            order += 1

    for table in doc.tables:
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows.append(' | '.join(cells))
        table_text = '\n'.join(rows).strip()
        if table_text:
            blocks.append({'type': 'text', 'content': table_text, 'path': None, 'order': order, 'page': None})
            order += 1

    for rel in doc.part.rels.values():
        if 'image' not in rel.reltype:
            continue
        img_idx += 1
        blob = rel.target_part.blob
        ctype = getattr(rel.target_part, 'content_type', '') or ''
        ext = '.png'
        if 'jpeg' in ctype or 'jpg' in ctype:
            ext = '.jpg'
        elif 'gif' in ctype:
            ext = '.gif'
        elif 'bmp' in ctype:
            ext = '.bmp'
        img_path = out_dir / f'docx_img_{img_idx}{ext}'
        img_path.write_bytes(blob)
        blocks.append({'type': 'image', 'content': None, 'path': str(img_path), 'order': order, 'page': None})
        order += 1

    blocks.sort(key=lambda item: item['order'])
    return blocks
