import tempfile
import unittest
from pathlib import Path

from docx import Document

from app.services.document_processor import blocks_to_plain_text, load_document


class DocumentProcessorTests(unittest.TestCase):
    def test_load_docx_blocks(self):
        with tempfile.TemporaryDirectory() as tmp:
            docx_path = Path(tmp) / 'sample.docx'
            image_dir = Path(tmp) / 'images'
            doc = Document()
            doc.add_paragraph('登录模块支持用户名密码登录')
            doc.add_paragraph('注册模块支持手机号注册')
            doc.save(str(docx_path))

            blocks = load_document(docx_path, image_dir)
            self.assertGreaterEqual(len(blocks), 2)
            self.assertTrue(all(block['type'] == 'text' for block in blocks))
            text = blocks_to_plain_text(blocks)
            self.assertIn('登录模块', text)


if __name__ == '__main__':
    unittest.main()
