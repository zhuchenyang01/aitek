from io import BytesIO

from openpyxl import Workbook, load_workbook

EVAL_HEADERS = ('query', 'relevant_chunk_ids', 'relevant_doc_ids', 'note')


def build_eval_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = 'queries'
    ws.append(list(EVAL_HEADERS))
    ws.append(['用户登录失败超过3次是否锁定账号？', '12,15', '3', '示例：请改成知识库里的真实分块 ID'])
    ws.append(['登录成功后如何签发 Token？', '12', '3', ''])
    ws.column_dimensions['A'].width = 48
    ws.column_dimensions['B'].width = 22
    ws.column_dimensions['C'].width = 18
    ws.column_dimensions['D'].width = 36

    readme = wb.create_sheet('readme')
    lines = [
        '使用说明',
        '1. 先在「AI 生成测试用例」页创建知识库并上传 pdf/docx。',
        '2. 点击「查看」拿到真实的分块 ID、文档 ID。',
        '3. 在 queries 表填写评测问题；相关 ID 用英文逗号分隔，例如 12,15。',
        '4. 不要把分块原文整段当作问题，否则召回率会虚高。',
        '5. 填好后回到「检索评测」页上传本文件，仍可再改，再点开始评测。',
        '',
        '列说明',
        'query：必填，评测问题',
        'relevant_chunk_ids：必填，相关分块 ID，逗号分隔',
        'relevant_doc_ids：可选，相关文档 ID，逗号分隔',
        'note：可选备注',
    ]
    for line in lines:
        readme.append([line])
    readme.column_dimensions['A'].width = 80

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parse_eval_excel(file_obj):
    wb = load_workbook(file_obj, read_only=True, data_only=True)
    try:
        ws = wb['queries'] if 'queries' in wb.sheetnames else wb.active
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()

    if not rows:
        raise ValueError('Excel 为空')

    header_index = _find_header_row(rows)
    if header_index is None:
        raise ValueError('未找到表头，请使用官方模板（含 query / relevant_chunk_ids）')

    header = [_cell_str(item).lower() for item in rows[header_index]]
    col = {name: header.index(name) for name in EVAL_HEADERS if name in header}
    if 'query' not in col:
        raise ValueError('缺少 query 列')

    queries = []
    for row in rows[header_index + 1:]:
        if row is None:
            continue
        query = _cell_str(_row_get(row, col['query']))
        if not query:
            continue
        chunk_ids = _parse_ids(_row_get(row, col.get('relevant_chunk_ids')))
        doc_ids = _parse_ids(_row_get(row, col.get('relevant_doc_ids')))
        note = _cell_str(_row_get(row, col.get('note')))
        queries.append(
            {
                'query': query,
                'relevant_chunk_ids': chunk_ids,
                'relevant_doc_ids': doc_ids,
                'note': note,
            }
        )
    if not queries:
        raise ValueError('未解析到有效评测问题')
    return queries


def _find_header_row(rows):
    for index, row in enumerate(rows):
        names = {_cell_str(item).lower() for item in (row or [])}
        if 'query' in names:
            return index
    return None


def _row_get(row, index):
    if index is None or row is None or index >= len(row):
        return None
    return row[index]


def _cell_str(value):
    if value is None:
        return ''
    return str(value).strip()


def _parse_ids(value):
    if value is None or value == '':
        return []
    if isinstance(value, float) and value.is_integer():
        return [int(value)]
    if isinstance(value, int):
        return [int(value)]
    text = str(value).strip().replace('，', ',')
    ids = []
    for part in text.split(','):
        part = part.strip()
        if not part:
            continue
        try:
            ids.append(int(float(part)))
        except (TypeError, ValueError) as exc:
            raise ValueError(f'无法解析 ID：{part}') from exc
    return ids
