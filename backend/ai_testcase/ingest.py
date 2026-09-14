import os
import tempfile
from pathlib import Path

from rag.pipeline.step4_chunk_and_embed import chunk_text

from rag.utils.vector_store import packed_vector
from system.llm import resolve_embedding_client

from .generator import _embedder
from .models import KnowledgeChunk, KnowledgeDocument

ALLOWED_EXTENSIONS = {'.pdf', '.docx'}


def ingest_document(knowledge_base, title, content, embedder=None, source_filename=''):
    text = (content or '').strip()
    if not text:
        raise ValueError('文档内容不能为空')

    document = KnowledgeDocument.objects.create(
        knowledge_base=knowledge_base,
        title=(title or '').strip(),
        source_filename=(source_filename or '').strip(),
        content=text,
    )
    chunks = chunk_text(text)
    if not chunks:
        return document

    embedder = embedder or resolve_embedding_client(knowledge_base.user) or _embedder()
    vectors = embedder.embed_batch(chunks)
    rows = []
    for index, (chunk, vector) in enumerate(zip(chunks, vectors)):
        rows.append(
            KnowledgeChunk(
                knowledge_base=knowledge_base,
                document=document,
                content=chunk,
                vector=packed_vector(vector),
                chunk_index=index,
            )
        )
    KnowledgeChunk.objects.bulk_create(rows)
    return document


def ingest_uploaded_file(knowledge_base, uploaded, title='', embedder=None):
    filename = (getattr(uploaded, 'name', None) or 'document').strip()
    suffix = Path(filename).suffix.lower()
    if suffix == '.doc':
        raise ValueError('暂仅支持 .docx，请先转换为 docx')
    if suffix not in ALLOWED_EXTENSIONS:
        raise ValueError('仅支持 pdf 或 docx 文件')

    tmp_path = None
    try:
        with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
            for chunk in uploaded.chunks():
                tmp.write(chunk)
            tmp_path = tmp.name
        text = extract_text_from_file(tmp_path, filename)
        display_title = (title or '').strip() or Path(filename).stem
        return ingest_document(
            knowledge_base,
            title=display_title,
            content=text,
            source_filename=filename,
            embedder=embedder,
        )
    finally:
        if tmp_path and os.path.exists(tmp_path):
            os.unlink(tmp_path)


def extract_text_from_file(path, filename=None):
    suffix = Path(filename or path).suffix.lower()
    if suffix == '.doc':
        raise ValueError('暂仅支持 .docx，请先转换为 docx')
    if suffix == '.pdf':
        text = _extract_pdf(path)
    elif suffix == '.docx':
        text = _extract_docx(path)
    else:
        raise ValueError('仅支持 pdf 或 docx 文件')
    if not (text or '').strip():
        raise ValueError('未能从文件中解析出文本')
    return text.strip()


def _extract_pdf(path):
    from pypdf import PdfReader

    reader = PdfReader(path)
    parts = []
    for page in reader.pages:
        text = (page.extract_text() or '').strip()
        if text:
            parts.append(text)
    return '\n\n'.join(parts)


def _extract_docx(path):
    from docx import Document

    doc = Document(path)
    parts = []
    for para in doc.paragraphs:
        text = (para.text or '').strip()
        if text:
            parts.append(text)
    for table in doc.tables:
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows.append(' | '.join(cells))
        table_text = '\n'.join(rows).strip()
        if table_text:
            parts.append(table_text)
    return '\n\n'.join(parts)
