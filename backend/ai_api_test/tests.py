import json
from io import BytesIO
from unittest.mock import MagicMock, patch

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from docx import Document
from rest_framework.authtoken.models import Token
from rest_framework.test import APIClient
from django.test import TestCase

from functional_test.models import FunctionalProject
from ai_api_test.models import ApiInterface, ApiInterfaceData, ApiTestCase, ApiTestRun


class ApiInterfaceApiTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='api_user', password='pass12345')
        self.token = Token.objects.create(user=self.user)
        self.client = APIClient()
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {self.token.key}')
        self.project = FunctionalProject.objects.create(user=self.user, name='贷款项目')

    def test_create_list_and_detail(self):
        created = self.client.post(
            '/api/ai-api-test/interfaces/',
            {
                'project_id': self.project.id,
                'name': '查询用户',
                'method': 'get',
                'url': 'https://api.example.com/users/{id}',
                'path': '/users/{id}',
                'version': 'v1',
                'headers': [{'key': 'Accept', 'value': 'application/json', 'enabled': True}],
                'query_params': [{'key': 'page', 'value': '1'}],
                'body_mode': 'none',
                'response_status': '200',
                'response_body': '{"ok":true}',
            },
            format='json',
        )
        self.assertEqual(created.status_code, 201, created.data)
        item_id = created.data['data']['id']
        self.assertTrue(created.data['data']['uid'])

        listed = self.client.get('/api/ai-api-test/interfaces/')
        self.assertEqual(listed.status_code, 200)
        payload = listed.data['data']
        self.assertEqual(payload['total'], 1)
        self.assertEqual(payload['page'], 1)
        self.assertEqual(payload['page_size'], 20)
        rows = payload['items']
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['name'], '查询用户')
        self.assertEqual(rows[0]['method'], 'GET')
        self.assertEqual(rows[0]['url'], 'https://api.example.com/users/{id}')
        self.assertNotIn('headers', rows[0])

        detail = self.client.get(f'/api/ai-api-test/interfaces/{item_id}/')
        self.assertEqual(detail.status_code, 200)
        data = detail.data['data']
        self.assertEqual(data['id'], item_id)
        self.assertEqual(data['headers'][0]['key'], 'Accept')
        self.assertEqual(data['response_body'], '{"ok":true}')
        self.assertEqual(data['project_name'], '贷款项目')

        updated = self.client.put(
            f'/api/ai-api-test/interfaces/{item_id}/',
            {
                'project_id': self.project.id,
                'name': '查询用户详情',
                'method': 'GET',
                'url': 'https://api.example.com/users/{id}',
                'version': 'v2',
                'description': '按 ID 查询',
                'headers': [{'key': 'Accept', 'value': 'application/json'}],
                'body_mode': 'none',
            },
            format='json',
        )
        self.assertEqual(updated.status_code, 200, updated.data)
        self.assertEqual(updated.data['data']['name'], '查询用户详情')
        self.assertEqual(updated.data['data']['version'], 'v2')
        self.assertEqual(ApiInterface.objects.get(pk=item_id).name, '查询用户详情')

        deleted = self.client.delete(f'/api/ai-api-test/interfaces/{item_id}/')
        self.assertEqual(deleted.status_code, 200, deleted.data)
        self.assertFalse(ApiInterface.objects.filter(pk=item_id).exists())

    def test_import_openapi(self):
        spec = {
            'openapi': '3.0.0',
            'info': {'title': 'Demo', 'version': '1.0.0'},
            'servers': [{'url': 'https://api.example.com'}],
            'paths': {
                '/pets': {
                    'get': {
                        'summary': '列出宠物',
                        'parameters': [
                            {'in': 'query', 'name': 'limit', 'schema': {'type': 'integer'}},
                        ],
                        'responses': {
                            '200': {
                                'description': 'ok',
                                'content': {
                                    'application/json': {
                                        'example': [{'id': 1, 'name': 'dog'}],
                                    }
                                },
                            }
                        },
                    },
                    'post': {
                        'summary': '创建宠物',
                        'requestBody': {
                            'content': {
                                'application/json': {
                                    'example': {'name': 'cat'},
                                }
                            }
                        },
                        'responses': {'201': {'description': 'created'}},
                    },
                }
            },
        }
        uploaded = self.client.post(
            '/api/ai-api-test/interfaces/import/',
            {
                'project_id': self.project.id,
                'version': 'v2',
                'file': SimpleUploadedFile(
                    'openapi.json',
                    json.dumps(spec).encode('utf-8'),
                    content_type='application/json',
                ),
            },
            format='multipart',
        )
        self.assertEqual(uploaded.status_code, 201, uploaded.data)
        self.assertEqual(uploaded.data['data']['count'], 2)
        self.assertEqual(ApiInterface.objects.filter(user=self.user).count(), 2)
        pet_list = ApiInterface.objects.get(name='列出宠物')
        self.assertEqual(pet_list.method, 'GET')
        self.assertEqual(pet_list.version, 'v2')
        self.assertEqual(pet_list.query_params[0]['key'], 'limit')
        self.assertIn('dog', pet_list.response_body)

    @patch('ai_api_test.views.resolve_llm_client')
    def test_import_docx_uses_llm(self, mock_resolve):
        mock_llm = MagicMock()
        mock_llm.chat.return_value = json.dumps(
            [
                {
                    'name': '登录',
                    'method': 'POST',
                    'url': '/api/login',
                    'path': '/api/login',
                    'description': '账号登录',
                    'body_mode': 'json',
                    'body_raw': '{"user":"a"}',
                    'response_status': '200',
                    'response_body': '{"token":"x"}',
                }
            ]
        )
        mock_resolve.return_value = mock_llm

        buf = BytesIO()
        doc = Document()
        doc.add_paragraph('接口：登录 POST /api/login')
        doc.add_paragraph('请求体 {"user":"a"} 响应 {"token":"x"}')
        doc.save(buf)
        uploaded = self.client.post(
            '/api/ai-api-test/interfaces/import/',
            {
                'project_id': self.project.id,
                'file': SimpleUploadedFile(
                    'apis.docx',
                    buf.getvalue(),
                    content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                ),
            },
            format='multipart',
        )
        self.assertEqual(uploaded.status_code, 201, uploaded.data)
        self.assertEqual(uploaded.data['data']['count'], 1)
        item = ApiInterface.objects.get(name='登录')
        self.assertEqual(item.method, 'POST')
        self.assertEqual(item.source_type, 'word')
        self.assertTrue(mock_llm.chat.called)


class ApiInterfaceDataApiTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='data_user', password='pass12345')
        self.token = Token.objects.create(user=self.user)
        self.client = APIClient()
        self.client.credentials(HTTP_AUTHORIZATION=f'Token {self.token.key}')
        self.project = FunctionalProject.objects.create(user=self.user, name='数据项目')
        self.interface = ApiInterface.objects.create(
            user=self.user,
            project=self.project,
            name='登录',
            method='POST',
            url='/api/login',
        )

    def test_dataset_crud(self):
        created = self.client.post(
            '/api/ai-api-test/datasets/',
            {
                'interface_id': self.interface.id,
                'description': '正确的账号和密码',
                'payload': {'username': 'user001', 'password': 'secret'},
            },
            format='json',
        )
        self.assertEqual(created.status_code, 201, created.data)
        pk = created.data['data']['id']
        self.assertEqual(created.data['data']['interface_id'], self.interface.id)

        listed = self.client.get('/api/ai-api-test/datasets/')
        self.assertEqual(listed.status_code, 200)
        self.assertEqual(listed.data['data']['total'], 1)

        updated = self.client.put(
            f'/api/ai-api-test/datasets/{pk}/',
            {
                'interface_id': self.interface.id,
                'description': '错误密码',
                'payload': {'username': 'user001', 'password': 'wrong'},
            },
            format='json',
        )
        self.assertEqual(updated.status_code, 200, updated.data)
        self.assertEqual(updated.data['data']['description'], '错误密码')

        deleted = self.client.delete(f'/api/ai-api-test/datasets/{pk}/')
        self.assertEqual(deleted.status_code, 200)
        self.assertEqual(ApiInterfaceData.objects.count(), 0)

    def test_dataset_rejects_unknown_interface(self):
        resp = self.client.post(
            '/api/ai-api-test/datasets/',
            {'interface_id': 99999, 'description': 'x', 'payload': {}},
            format='json',
        )
        self.assertEqual(resp.status_code, 400)


class ApiTestCaseApiTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='case_user', password='pass12345')
        self.token = Token.objects.create(user=self.user)
        self.client = APIClient()
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {self.token.key}')
        self.project = FunctionalProject.objects.create(user=self.user, name='用例项目')
        self.interface = ApiInterface.objects.create(
            user=self.user,
            project=self.project,
            name='登录',
            method='POST',
            url='/api/login',
            body_raw='{"username":"","password":""}',
        )

    @patch('ai_api_test.views.resolve_llm_client')
    def test_generate_and_crud_cases(self, mock_resolve):
        mock_llm = MagicMock()
        mock_llm.chat.return_value = json.dumps(
            [
                {
                    'name': '正确账号登录成功',
                    'priority': 'P0',
                    'case_type': '正向',
                    'description': '使用正确账号密码登录',
                    'request_body': {'username': 'user001', 'password': 'ok'},
                    'expected_status': '200',
                    'assertions': [{'name': 'status', 'expected': '200'}],
                }
            ]
        )
        mock_resolve.return_value = mock_llm

        generated = self.client.post(f'/api/ai-api-test/interfaces/{self.interface.id}/generate-cases/')
        self.assertEqual(generated.status_code, 201, generated.data)
        self.assertEqual(generated.data['data']['count'], 1)
        self.assertEqual(ApiTestCase.objects.count(), 1)

        listed = self.client.get('/api/ai-api-test/cases/')
        self.assertEqual(listed.status_code, 200)
        self.assertEqual(listed.data['data']['total'], 1)
        pk = listed.data['data']['items'][0]['id']

        updated = self.client.put(
            f'/api/ai-api-test/cases/{pk}/',
            {
                'interface_id': self.interface.id,
                'name': '正确账号登录成功-改',
                'priority': 'P1',
                'description': '已编辑',
                'expected_status': '200',
            },
            format='json',
        )
        self.assertEqual(updated.status_code, 200, updated.data)
        self.assertEqual(updated.data['data']['name'], '正确账号登录成功-改')

        deleted = self.client.delete(f'/api/ai-api-test/cases/{pk}/')
        self.assertEqual(deleted.status_code, 200)
        self.assertEqual(ApiTestCase.objects.count(), 0)

    @patch('ai_api_test.views.resolve_llm_client')
    def test_generate_cases_stream(self, mock_resolve):
        mock_llm = MagicMock()
        mock_llm.iter_chat_events.return_value = iter(
            [
                {'type': 'thinking', 'content': '先覆盖正向登录'},
                {
                    'type': 'answer',
                    'content': '【思考】覆盖正向登录\n【用例】'
                    + json.dumps(
                        [
                            {
                                'name': '流式生成登录成功',
                                'priority': 'P0',
                                'case_type': '正向',
                                'expected_status': '200',
                                'assertions': [{'name': 'status', 'expected': '200'}],
                            }
                        ],
                        ensure_ascii=False,
                    ),
                },
            ]
        )
        mock_resolve.return_value = mock_llm
        resp = self.client.post(f'/api/ai-api-test/interfaces/{self.interface.id}/generate-cases/stream/')
        self.assertEqual(resp.status_code, 200)
        body = b''.join(resp.streaming_content).decode('utf-8')
        self.assertIn('已生成 1 条测试用例', body)
        self.assertEqual(ApiTestCase.objects.filter(name='流式生成登录成功').count(), 1)

    @patch('ai_api_test.views.run_api_case')
    def test_run_case_proxies_service(self, mock_run):
        interface = ApiInterface.objects.create(
            user=self.user,
            project=self.project,
            name='登录',
            method='POST',
            url='https://example.com/login',
        )
        case = ApiTestCase.objects.create(
            user=self.user,
            interface=interface,
            name='登录成功',
            expected_status='200',
            request_body='{"username":"a"}',
        )
        mock_run.return_value = {
            'code': 0,
            'msg': '接口测试执行成功',
            'data': {
                'passed': True,
                'run_id': 'run-1',
                'http': {'status_code': 200},
                'result_logs': [
                    {
                        'event': 'addSuccess',
                        'outcome': 'success',
                        'test_name': 'test_api',
                        'message': 'passed',
                        'traceback': '',
                        'screenshot_path': '/tmp/a.png',
                        'http': {'status_code': 200},
                    }
                ],
            },
        }
        resp = self.client.post(f'/api/ai-api-test/cases/{case.id}/run/')
        self.assertEqual(resp.status_code, 200, resp.data)
        self.assertTrue(resp.data['data']['passed'])
        self.assertEqual(case.run_logs.count(), 1)
        self.assertEqual(ApiTestRun.objects.filter(case=case).count(), 0)
        listed = self.client.get('/api/ai-api-test/runs/')
        self.assertEqual(listed.status_code, 200)
        self.assertEqual(listed.data['data']['total'], 0)

    @patch('ai_api_test.suite_views.run_api_suite')
    def test_suite_run_and_report(self, mock_suite):
        interface = ApiInterface.objects.create(
            user=self.user,
            project=self.project,
            name='登录',
            method='POST',
            url='https://example.com/login',
        )
        case = ApiTestCase.objects.create(
            user=self.user,
            interface=interface,
            name='登录成功',
            expected_status='200',
        )
        created = self.client.post(
            '/api/ai-api-test/suites/',
            {
                'project_id': self.project.id,
                'name': '登录链路',
                'steps': [{'case_id': case.id, 'extractors': [{'name': 'token', 'path': 'data.token'}]}],
            },
            format='json',
        )
        self.assertEqual(created.status_code, 201, created.data)
        suite_id = created.data['data']['id']
        mock_suite.return_value = {
            'code': 0,
            'msg': '套件执行成功',
            'data': {
                'passed': True,
                'run_id': 'suite-1',
                'total': 1,
                'passed_count': 1,
                'failed_count': 0,
                'skipped_count': 0,
                'steps': [
                    {
                        'step_index': 1,
                        'case_id': case.id,
                        'case_name': '登录成功',
                        'passed': True,
                        'http': {'status_code': 200},
                        'extracts': {'token': 'abc'},
                        'assertion_logs': [],
                        'result_logs': [],
                        'screenshots': [],
                        'unittest_output': '',
                        'failures': [],
                    }
                ],
            },
        }
        ran = self.client.post(f'/api/ai-api-test/suites/{suite_id}/run/')
        self.assertEqual(ran.status_code, 200, ran.data)
        self.assertTrue(ran.data['data']['passed'])
        run = ApiTestRun.objects.get(suite_id=suite_id)
        detail = self.client.get(f'/api/ai-api-test/runs/{run.id}/')
        self.assertEqual(detail.status_code, 200)
        self.assertEqual(detail.data['data']['steps'][0]['extracts']['token'], 'abc')
        self.assertIn('request', detail.data['data']['steps'][0])
        self.assertIn('response', detail.data['data']['steps'][0])
